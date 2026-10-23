import base64
from datetime import datetime
import io
import json
import anthropic
import openpyxl
from openpyxl.styles import Alignment
import streamlit as st

# Titel der Anwendung
st.title("🦷 Dental-Rechnungs-Generator")

# 1. Manuelle Eingabe der Rechnungsnummer (z.B. 191)
st.subheader("1. Rechnungsnummer eingeben")
eingabe_nr = st.text_input(
    "Rechnungsnummer (nur die Nummer, z. B. 191):", placeholder="191"
)

# 2. Laborauftrag fotografieren / hochladen
st.subheader("2. Laborauftrag fotografieren / hochladen")
uploaded_file = st.file_uploader(
    "Auftragszettel als Bild hochladen", type=["jpg", "jpeg", "png"]
)

if uploaded_file is not None:
  st.image(
      uploaded_file, caption="Aufgenommener Zettel", use_column_width=True
  )

  if st.button("🚀 Rechnung automatisch befüllen"):
    with st.spinner(
        "Claude liest den Zettel aus und befüllt deine Vorlage..."
    ):
      try:
        # Bild in Base64 umwandeln
        image_bytes = uploaded_file.getvalue()
        base64_image = base64.b64encode(image_bytes).decode("utf-8")

        # Anthropic Client initialisieren
        client = anthropic.Anthropic(api_key=st.secrets["ANTHROPIC_API_KEY"])

        # Prompt für Claude zur Extraktion
        prompt_text = """Analysiere diesen zahnärztlichen Laborauftrag exakt. 
Antworte ausschließlich im folgenden JSON-Format (kein Markdown drumherum, nur das reine JSON):
{
  "patientenname": "...",
  "auftragsnummer": "..."
}"""

        message = client.messages.create(
            model="claude-sonnet-5-5",
            max_tokens=1500,
            messages=[
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "image",
                            "source": {
                                "type": "base64",
                                "media_type": "image/jpeg",
                                "data": base64_image,
                            },
                        },
                        {"type": "text", "text": prompt_text},
                    ],
                }
            ],
        )

        # Antwort sicher extrahieren
        result_text = ""
        for block in message.content:
          if hasattr(block, "text"):
            result_text += block.text

        # JSON parsen
        json_start = result_text.find("{")
        json_end = result_text.rfind("}") + 1
        json_data = json.loads(result_text[json_start:json_end])

        patientenname = json_data.get("patientenname", "Unbekannt")
        auftragsnummer = json_data.get("auftragsnummer", "Unbekannt")

        # Rechnungsnummer zusammenbauen (z. B. 26.191 aus Eingabe 191)
        sauber_nr = eingabe_nr.strip()
        rechnungsnummer_voll = f"26.{sauber_nr}" if sauber_nr else "26.000"

        # Feste Excel-Vorlage aus dem Repository laden
        wb = openpyxl.load_workbook("Vorlage.xlsx")
        ws = wb.active  # Nimmt das aktive Tabellenblatt

        # Zellen exakt nach deiner Vorgabe belegen:
        ws["G9"] = datetime.now().strftime("%d.%m.%Y")
        ws["G9"].alignment = Alignment(
            horizontal="right"
        )  # Datum rechtsbündig

        ws["G11"] = (
            rechnungsnummer_voll  # Rechnungsnummer (z.B. 26.191) in G11
        )
        ws["G12"] = patientenname  # Patientenname vom Zettel in G12
        ws["G16"] = auftragsnummer  # XML-Nummer / Auftragsnummer in G16

        # In den Arbeitsspeicher speichern
        output = io.BytesIO()
        wb.save(output)
        excel_data = output.getvalue()

        st.success("Rechnung wurde erfolgreich befüllt!")

        # Download-Button mit dem gewünschten Dateinamen (z.B. 26.191_Rechnung.xlsx)
        st.download_button(
            label="📥 Befüllte Rechnung herunterladen (.xlsx)",
            data=excel_data,
            file_name=f"{rechnungsnummer_voll}_Rechnung.xlsx",
            mime=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
        )

      except Exception as e:
        st.error(f"Fehler bei der Verarbeitung: {e}")
else:
  st.info(
      "💡 Gib deine Rechnungsnummer (z. B. 191) ein und lade ein Foto des"
      " Auftragszettels hoch."
  )
